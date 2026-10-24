"""The Excel sheet of voters."""
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

import export
import tg
from conftest import user


def make_voters(n, cap=20):
    out = []
    for i in range(1, n + 1):
        v = tg.Voter(i, user(i, username=i % 3 != 0), datetime(2026, 10, 24, 10, i % 60, tzinfo=timezone.utc))
        v.status = "Added" if i <= cap else f"Not in first {cap}"
        v.note = "" if i % 4 else "Privacy settings block being added"
        out.append(v)
    return out


def test_sheet_contents(fake):
    name = export.write_run("28 Oct Recre @ B+ Aperia", make_voters(25), 20)
    ws = load_workbook(export.EXPORT_DIR / name).worksheets[0]
    assert [c.value for c in ws[1]] == export.HEADERS
    assert ws.max_row == 26
    assert ws["A2"].value == 1 and ws["C2"].value == "@climber1" and ws["C4"].value == "(no handle)"
    assert ws["F21"].value == "Added" and ws["F22"].value == "Not in first 20"
    assert ws["A22"].fill.fgColor.rgb.endswith("EEEEEE")


def test_every_run_adds_a_tab_newest_first(fake):
    for _ in range(40):
        name = export.write_run("Recre", make_voters(3), 20)
    wb = load_workbook(export.EXPORT_DIR / name)
    assert len(wb.sheetnames) == 40
    assert len(set(wb.sheetnames)) == 40
    assert all(len(s) <= 31 for s in wb.sheetnames)
    assert sum(ws.sheet_view.tabSelected for ws in wb.worksheets) == 1


@pytest.mark.parametrize("title", [
    'a/b\\c:d*e?f"g<h>i|j', "CON", "nul", "Recre 🧗 @ GU", "...", "trailing dots...", "  ", "x" * 128,
])
def test_awkward_titles_make_valid_files(fake, title):
    name = export.write_run(title, make_voters(2), 20)
    assert (export.EXPORT_DIR / name).exists()
    stem = name.rsplit(".", 1)[0]
    assert stem.upper() not in {"CON", "NUL", "PRN", "AUX"}
    assert not stem.endswith((".", " "))


def test_empty_run(fake):
    name = export.write_run("Empty", [], 20)
    assert load_workbook(export.EXPORT_DIR / name).worksheets[0].max_row == 1
