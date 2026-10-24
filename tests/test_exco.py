"""Exco: always added if they vote, and never use up one of the spots."""
from datetime import timedelta

import pytest
from openpyxl import load_workbook

import app as webapp
import export
import store
import tg
from conftest import CHAT, OPT_A, OPT_B, POLL_DATE, POLL_MSG, run_option, statuses, user


def set_exco(*handles):
    store.update_config(exco=list(handles))


def voters():
    return tg.run(tg.get_voters(CHAT, POLL_MSG, OPT_A.hex()))


# ---------------------------------------------------------------- reading the list

@pytest.mark.parametrize("text,good,bad", [
    ("@Alice\n@bob_climbs", ["alice", "bob_climbs"], []),
    ("alice, bob_climbs; carol_x", ["alice", "bob_climbs", "carol_x"], []),
    ("https://t.me/Alice t.me/bob_climbs telegram.me/carol_x", ["alice", "bob_climbs", "carol_x"], []),
    ("@alice\n@ALICE\n alice ", ["alice"], []),
    ("@ab @1abc @has-dash @" + "x" * 40, [], ["@ab", "@1abc", "@has-dash", "@" + "x" * 40]),
    ("", [], []),
    ("\n\n  \n", [], []),
])
def test_parse_handles(text, good, bad):
    assert store.parse_handles(text) == (good, bad)


# ---------------------------------------------------------------- counting

def test_exco_not_numbered_and_dont_use_spots(fake):
    fake.vote(OPT_A, range(1, 8))
    set_exco("climber2", "climber5")
    got = voters()
    assert [(v.id, v.position, v.exco) for v in got] == [
        (1, 1, False), (2, None, True), (3, 2, False), (4, 3, False),
        (5, None, True), (6, 4, False), (7, 5, False)]


def test_handle_match_ignores_case(fake):
    fake.users[1] = user(1)
    fake.users[1].username = "Climber_ONE"
    fake.vote(OPT_A, [1])
    set_exco("climber_one")
    assert voters()[0].exco


def test_voter_without_handle_is_never_exco(fake):
    fake.users[1] = user(1, username=False)
    fake.vote(OPT_A, [1])
    set_exco("climber1")
    assert not voters()[0].exco


def test_exco_voting_late_still_added(fake):
    fake.vote(OPT_A, range(1, 31))
    set_exco("climber30")
    s = statuses(run_option(fake))
    assert s[30] == "Added"
    assert sum(x == "Added" for x in s.values()) == 21  # 20 spots + exco
    assert s[21] == "Not in first 20"


def test_exco_voting_early_frees_a_spot(fake):
    fake.vote(OPT_A, range(1, 26))
    set_exco("climber1", "climber2", "climber3")
    s = statuses(run_option(fake))
    assert all(s[i] == "Added" for i in range(1, 24))  # 3 exco + the next 20
    assert s[24] == s[25] == "Not in first 20"


def test_exco_who_did_not_vote_are_not_added(fake):
    fake.vote(OPT_A, [1, 2])
    fake.users[99] = user(99)
    set_exco("climber99")
    run_option(fake)
    gid = next(iter(fake.groups))
    assert 99 not in fake.members[gid]


def test_exco_voting_for_other_option_only(fake):
    fake.vote(OPT_A, [1])
    fake.vote(OPT_B, [2])
    set_exco("climber2")
    run_option(fake)
    assert 2 not in fake.members[next(iter(fake.groups))]


def test_exco_with_privacy_block_gets_link(fake):
    fake.vote(OPT_A, range(1, 25))
    fake.privacy_block = {24}
    set_exco("climber24")
    assert statuses(run_option(fake))[24] == "Invite link sent"


def test_exco_added_on_rerun_after_list_changes(fake):
    fake.vote(OPT_A, range(1, 26))
    run_option(fake)
    set_exco("climber25")
    s = statuses(run_option(fake))
    assert s[25] == "Added"


def test_all_voters_exco(fake):
    fake.vote(OPT_A, range(1, 6))
    set_exco(*[f"climber{i}" for i in range(1, 6)])
    r = run_option(fake, cap=2)
    assert all(v.position is None for v in r["voters"])
    assert all(s == "Added" for s in statuses(r).values())


def test_excel_marks_exco(fake):
    fake.vote(OPT_A, range(1, 4))
    set_exco("climber2")
    r = run_option(fake)
    ws = load_workbook(export.EXPORT_DIR / export.write_run(r["title"], r["voters"], 20)).worksheets[0]
    assert [ws.cell(row=i, column=1).value for i in range(2, 5)] == [1, "Exco", 2]


# ---------------------------------------------------------------- pages

@pytest.fixture
def client(fake):
    store.update_config(api_id="123", api_hash="a" * 32)
    webapp.app.config["TESTING"] = False
    return webapp.app.test_client()


def test_exco_page_save_and_show(client):
    r = client.post("/exco", data={"handles": "@Alice\nt.me/bob_climbs\n@x\nalice"})
    html = client.get(r.headers["Location"]).get_data(as_text=True)
    assert store.exco_handles() == {"alice", "bob_climbs"}
    assert "Saved 2 exco handles" in html and "skipped: @x" in html
    assert "@alice\n@bob_climbs" in html


def test_exco_page_clear(client):
    set_exco("alice")
    client.post("/exco", data={"handles": ""})
    assert store.exco_handles() == set()


def test_option_page_shows_exco(client, fake):
    fake.vote(OPT_A, range(1, 24))
    set_exco("climber1", "climber2")
    html = client.get(f"/option?chat={CHAT}&msg={POLL_MSG}&opt={OPT_A.hex()}").get_data(as_text=True)
    assert html.count('<span class="tag exco">exco</span>') == 2
    assert "2 exco" in html
    assert "cut-off: first 20" in html
    assert "Create group and add 22 people" in html  # 2 exco + 20


def test_option_page_no_cutoff_when_exco_make_it_fit(client, fake):
    fake.vote(OPT_A, range(1, 22))
    set_exco("climber1")
    html = client.get(f"/option?chat={CHAT}&msg={POLL_MSG}&opt={OPT_A.hex()}").get_data(as_text=True)
    assert "cut-off" not in html and "Create group and add 21 people" in html


def test_nav_has_exco_button(client):
    html = client.get("/exco").get_data(as_text=True)
    assert 'class="button nav" href="/exco"' in html
