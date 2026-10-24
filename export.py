"""Excel sheet of voters: one workbook per group, one tab per run."""
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from store import EXPORT_DIR

HEADERS = ["Position", "Name", "Telegram handle", "User ID", "Voted at", "Status", "Notes"]
WIDTHS = [10, 26, 22, 14, 20, 34, 50]


_RESERVED = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}


def safe_filename(title):
    """A file name Windows and Mac both accept."""
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "", title).strip().rstrip(". ")
    if not name or name.upper() in _RESERVED:
        name = f"recre {name}".strip()
    return name


def write_run(title, voters, cap):
    path = EXPORT_DIR / f"{safe_filename(title)}.xlsx"
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    ws = wb.create_sheet(datetime.now().strftime("Run %d %b %H%M %S")[:31], 0)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    grey = PatternFill("solid", fgColor="EEEEEE")
    for v in voters:
        ws.append(["Exco" if v.exco else v.position, v.name, v.handle or "(no handle)", v.id,
                   v.voted_at.strftime("%d %b %Y %H:%M:%S"), v.status, getattr(v, "note", "")])
        if not v.gets_in(cap):
            for cell in ws[ws.max_row]:
                cell.fill = grey
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    for sheet in wb.worksheets:  # only the new tab selected, else Excel groups them
        sheet.sheet_view.tabSelected = sheet is ws
    wb.active = 0
    wb.save(path)
    return path.name
